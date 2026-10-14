"""Import and export records as CSV or Excel.

Export uses the table's own columns (with their formatting) and its current
search/filters. Import uses an :class:`Importer` that lists the columns::

    class ProductImporter(Importer):
        model = Product
        columns = [
            ImportColumn("name").required(),
            ImportColumn("price").numeric(),
            ImportColumn("category").relationship("category", "name"),
        ]

    table.header_actions([ImportAction(ProductImporter), ExportAction()])

An export action can list its own :class:`ExportColumn` s instead of using the
table's columns::

    ExportAction(columns=[ExportColumn("sku"), ExportColumn("price", format=lambda state: f"{state:.2f}")])

Big exports can run in the background. The user gets a "ready to download"
notification, and every export is kept in their export history
(``<panel>/_tw/exports``)::

    ExportAction().background()                # always
    ExportAction().background(min_rows=5000)   # only when 5,000 rows or more match
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import tempfile
from decimal import Decimal, InvalidOperation
from typing import TYPE_CHECKING, Any, Callable, ClassVar
from urllib.parse import parse_qsl, urlencode

from .actions.action import Action, BulkAction
from .i18n import translate as __
from .support.component import headline
from .support.evaluate import call

if TYPE_CHECKING:  # pragma: no cover
    from .context import Context
    from .hosts import Host


# ---------------------------------------------------------------------- export
class ExportColumn:
    """One column of an export file: a (dotted) attribute path, a header and an optional ``format(state, record)``."""

    def __init__(self, name: str, label: str | None = None, format: Callable | None = None) -> None:
        self.name = name
        self.label = label or headline(name.replace(".", " "))
        self.format = format

    def get_label(self) -> str:
        return __(self.label)

    def get_value(self, record: Any) -> Any:
        from .tables.columns import read_path

        state = read_path(record, self.name)
        if isinstance(state, list):
            state = ", ".join(str(s) for s in state)
        return call(self.format, state=state, record=record) if self.format else state


def export_rows(ctx: "Context", host: "Host", params: Any, keys: list[str] | None, columns: list[str] | None,
                export_columns: list[ExportColumn] | None = None):
    """Yield the header row and then one row per record.

    Uses the table's columns, or ``export_columns`` when the export action lists its own.
    """
    from .tables.columns import Column, ImageColumn, ToggleColumn, read_path
    from .tables.table import Table

    table: Table = host.get_table(ctx)
    table.bind(ctx, host, params=params)
    cols: list = [c for c in table._columns if isinstance(c, Column) and not isinstance(c, ImageColumn)]
    if export_columns:
        cols = list(export_columns)
    if columns:
        cols = [c for c in cols if c.name in columns]
    yield [c.get_label() for c in cols]
    if keys is not None:
        records = host.find_records(ctx, keys)
    else:
        query = table._eager_loads(table.sorted_query(table.filtered_query()))
        records = ctx.db.scalars(query).unique().all()
    table._page_records = list(records)  # ``counts()`` columns count every exported row in one query
    ev = table.ev()
    for record in records:
        if export_columns:
            yield [_plain(c.get_value(record)) for c in cols]
            continue
        row = []
        for c in cols:
            state = c.get_state(record, {**ev, "record": record}) if c._state else read_path(record, c.name)
            if isinstance(c, ToggleColumn):
                value = bool(state)
            elif isinstance(state, list):
                value = ", ".join(str(c.format_value(s, record, {**ev, "state": s}) or "") for s in state)
            else:
                value = c.format_value(state, record, {**ev, "state": state, "record": record})
            row.append(_plain(value))
        yield row


def _plain(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "__html__"):
        import re

        return re.sub(r"<[^>]+>", "", str(value))
    return value


def write_csv(rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in rows:
        writer.writerow([_csv_safe(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _csv_safe(value: Any) -> Any:
    """Stop spreadsheet formula injection (=, +, -, @ at the start)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def write_xlsx(rows) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel export needs openpyxl: pip install 'tungsten-admin[excel]'") from exc
    wb = Workbook()
    ws = wb.active
    for i, row in enumerate(rows):
        ws.append([_csv_safe(v) for v in row])
        if i == 0:
            for cell in ws[1]:
                cell.font = Font(bold=True)
    for column in ws.columns:
        width = max((len(str(c.value or "")) for c in column), default=10)
        ws.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def write_file(rows, fmt: str, out: Any) -> int:
    """Write ``rows`` (header first) to the binary file ``out`` as csv or xlsx. Returns the data row count."""
    count = -1
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Excel export needs openpyxl: pip install 'tungsten-admin[excel]'") from exc
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        for row in rows:
            count += 1
            if count == 0:
                header = []
                for v in row:
                    cell = WriteOnlyCell(ws, value=_csv_safe(v))
                    cell.font = Font(bold=True)
                    header.append(cell)
                ws.append(header)
            else:
                ws.append([_csv_safe(v) for v in row])
        wb.save(out)
        return max(count, 0)
    text = io.TextIOWrapper(out, encoding="utf-8", newline="", write_through=True)
    text.write("\ufeff")
    writer = csv.writer(text)
    for row in rows:
        count += 1
        writer.writerow([_csv_safe(v) for v in row])
    text.flush()
    text.detach()
    return max(count, 0)


def export_name(host: "Host") -> str:
    return re.sub(r"[^a-z0-9-]+", "-", (host.title() or "export").lower()).strip("-") or "export"


def record_export(ctx: "Context", host: "Host", fmt: str, rows: int, url: str) -> None:
    """Add a finished export to the user's export history (``url`` downloads it again with fresh data)."""
    from .models import Job

    if ctx.user is None:
        return
    panel = ctx.panel
    ctx.db.add(Job(
        kind="export", name=str(host.title() or "Export")[:255], status="done", processed=rows, total=rows,
        user_id=panel.auth.user_id(ctx.user),
        tenant_key=panel.tenancy.key_of(ctx.tenant) if ctx.tenant is not None and panel.tenancy.enabled else None,
        payload={"host": host.key, "format": fmt, "rerun": url},
        started_at=dt.datetime.now(), finished_at=dt.datetime.now(),
    ))
    ctx.db.commit()


#: background export files are deleted after this many days (their history entry stays)
EXPORT_FILE_DAYS = 7


def prune_exports(db: Any, storage: Any, days: int | None = None) -> None:
    """Delete background export files older than ``days`` (``EXPORT_FILE_DAYS``)."""
    from sqlalchemy import select

    from .models import Job

    cutoff = dt.datetime.now() - dt.timedelta(days=EXPORT_FILE_DAYS if days is None else days)
    old = db.scalars(select(Job).where(Job.kind == "export", Job.status == "done", Job.finished_at < cutoff)).all()
    for job in old:
        path = (job.result or {}).get("file")
        if not path:
            continue
        storage.delete(path)
        job.result = {**job.result, "file": None}
        job.status = "expired"
    if old:
        db.commit()


def run_export_job(job: Any):
    """Job handler: write the export file, save it to storage and notify the user."""
    from starlette.datastructures import QueryParams

    from .notifications import Notification
    from .routes import Forbidden, Routes

    ctx = job.ctx
    p = job.payload
    if ctx.user is None and ctx.panel.auth.enabled:
        raise PermissionError("The user who started this export no longer exists.")
    host = Routes(ctx.panel).resolve_host(ctx, p.get("host"))
    params = QueryParams(p.get("query", ""))
    keys = p.get("keys") or None
    action = find_offered_action(ctx, host, ExportBulkAction if keys else ExportAction, p.get("action"),
                                 bulk=bool(keys))
    if action is None:
        raise Forbidden()
    fmt = "xlsx" if p.get("format") == "xlsx" else "csv"
    total = len(keys) if keys else None
    if total is None:
        try:
            from sqlalchemy import func, select

            table = host.get_table(ctx)
            table.bind(ctx, host, params=params)
            total = ctx.db.scalar(select(func.count()).select_from(table.filtered_query().subquery()))
        except Exception:  # noqa: BLE001 - the count only feeds the progress bar
            total = None
    job.progress(0, total=total)

    def counted(rows):
        for i, row in enumerate(rows):
            yield row
            if i and i % 1000 == 0:
                job.progress(i)

    rows = export_rows(ctx, host, params, keys, p.get("columns") or None, action.export_columns)
    filename = f"{export_name(host)}-{dt.datetime.now():%Y%m%d-%H%M}.{fmt}"
    with tempfile.TemporaryFile() as tmp:
        count = write_file(counted(rows), fmt, tmp)
        yield
        tmp.seek(0)
        stored = ctx.panel.storage.save(tmp, filename, "exports")
    job.job.result = {"file": stored, "filename": filename, "format": fmt}
    job.progress(count, total=count)
    url = ctx.panel.url("_tw", "exports", job.id, "download")
    if ctx.panel.database_notifications and job.job.user_id is not None:
        Notification("Your export is ready").body(
            f"{job.job.name}: {count} rows ({fmt.upper()})").success().icon("download").action(
            "Download", url).send_to_database(job.job.user_id, job.db)
    prune_exports(job.db, ctx.panel.storage)
    return {"file": stored, "filename": filename, "format": fmt, "rows": count}


def find_offered_action(ctx: "Context", host: "Host", kind: type, name: str | None = None, bulk: bool = False):
    """The first ``kind`` action that the host's table (bulk or header actions) or page header offers
    and the user may run, or None. Endpoints use it so they only do what the UI offers."""
    from .actions.action import flatten_actions

    table = host.get_table(ctx)
    pool = list(table._bulk_actions if bulk else table._header_actions) if table is not None else []
    if not bulk:
        getter = getattr(host, "all_page_actions", None) or host.page_actions
        pool += list(getter(ctx, None) or [])
    for a in flatten_actions(pool):
        if isinstance(a, kind) and (not name or a.name == name) and a.is_available(host, ctx):
            return a
    return None


class ExportAction(Action):
    """Download the table (current search & filters) as CSV or Excel.

    ``columns`` (a list of :class:`ExportColumn`) replaces the table's columns in the file.
    """

    def __init__(self, name: str = "export", formats: tuple[str, ...] = ("csv", "xlsx"),
                 columns: list[ExportColumn] | None = None) -> None:
        super().__init__(name)
        self._label = "Export"
        self._icon = "download"
        self._color = "gray"
        self._formats = formats
        self.export_columns = list(columns) if columns else None
        self._modal_heading = "Export records"
        self._modal_description = "Download the records that match your current search and filters."
        self._modal_submit_label = "Download"
        self._modal_icon = "download"
        self._modal_width = "md"
        self._form = True
        self._style = "button"
        self._background: Any = False
        self._min_rows: int | None = None

    def background(self, condition: Any = True, min_rows: int | None = None) -> "ExportAction":
        """Build the file in the background and notify the user when it is ready to download.

        ``min_rows`` only does so when at least that many rows match (smaller exports download at once)."""
        self._background = condition
        self._min_rows = min_rows
        return self

    def build_form(self, ctx, host, record=None, records=None):  # type: ignore[override]
        from markupsafe import Markup

        from .forms import CheckboxList, Hidden, Placeholder, Radio
        from .forms.form import Form
        from .tables.columns import ImageColumn

        if self.export_columns:
            cols = {c.name: c.get_label() for c in self.export_columns}
        else:
            table = host.get_table(ctx)
            cols = {c.name: c.get_label() for c in table._columns if not isinstance(c, ImageColumn)}
        labels = {"csv": "CSV (.csv)", "xlsx": "Excel (.xlsx)"}
        return Form().columns(1).schema([
            Radio("format").options({f: labels.get(f, f) for f in self._formats}).default(self._formats[0])
            .inline().required(),
            CheckboxList("columns").options(cols).default(list(cols)).columns(2).bulk_toggleable().required(),
            Hidden("state"),
            Placeholder("history").hidden_label().content(Markup(
                '<a href="{}" class="text-sm font-medium text-primary-600 hover:text-primary-500">{}</a>').format(
                ctx.url("_tw", "exports"), __("Export history"))),
        ])

    def fill(self, form, ctx, record=None, records=None):  # type: ignore[override]
        form.fill(None)
        params = [(k, v) for k, v in ctx.request.query_params.multi_items() if not k.startswith("_tw_") and k != "records"]
        form.set("state", urlencode(params))

    def run(self, ctx, host, record=None, records=None, data=None, form=None):  # type: ignore[override]
        data = data or {}
        state = parse_qsl(data.get("state") or "")
        query = [(k, v) for k, v in state if k not in ("host", "action", "format", "columns", "keys")]
        fmt = data.get("format", "csv")
        if self._in_background(ctx, host, query):
            start_export(ctx, host, self.name, fmt, query, data.get("columns") or [], None)
            return None
        query += [("host", host.key), ("action", self.name), ("format", fmt)]
        query += [("columns", c) for c in data.get("columns") or []]
        ctx.redirect(ctx.url("_tw", "export") + "?" + urlencode(query))

    def _in_background(self, ctx, host, query) -> bool:
        from .support.evaluate import evaluate

        if not evaluate(self._background, **self.ev(ctx, host=host)):
            return False
        if not self._min_rows:
            return True
        from starlette.datastructures import QueryParams

        table = host.get_table(ctx)
        table.bind(ctx, host, params=QueryParams(urlencode(query)))
        from sqlalchemy import func, select

        count = ctx.db.scalar(select(func.count()).select_from(table.filtered_query().subquery())) or 0
        return count >= self._min_rows


def start_export(ctx: "Context", host: "Host", action_name: str, fmt: str, query: list, columns: list,
                 keys: list | None) -> None:
    """Queue a background export and tell the user it has started."""
    from .notifications import Notification

    payload = {"host": host.key, "action": action_name, "format": fmt, "query": urlencode(query),
               "columns": list(columns), "keys": list(keys) if keys else None}
    ctx.panel.jobs.dispatch(ctx, "export", payload, name=str(host.title() or "Export"),
                            total=len(keys) if keys else None)
    Notification("Export started").body("You will get a notification when the file is ready to download.").info(
    ).action("Export history", ctx.url("_tw", "exports"), "gray").send(ctx)


class ExportBulkAction(BulkAction):
    """Download only the selected rows. ``columns`` works like :class:`ExportAction`'s."""

    def __init__(self, name: str = "export", format: str = "csv", columns: list[ExportColumn] | None = None) -> None:
        super().__init__(name)
        self._label = "Export"
        self._icon = "download"
        self._format = format
        self.export_columns = list(columns) if columns else None
        self._deselect_after = False

    def background(self, condition: Any = True, chunk_size: int | None = None) -> "ExportBulkAction":
        """Build the file in the background and notify the user when it is ready to download.
        ``condition`` may be a closure, e.g. ``lambda keys: len(keys) > 1000``."""
        self._export_background = condition
        return self

    def run(self, ctx, host, record=None, records=None, data=None, form=None):  # type: ignore[override]
        from .support.evaluate import evaluate

        keys = [host.record_key(r) for r in records or []]
        if evaluate(getattr(self, "_export_background", False), **{**self.ev(ctx), "keys": keys}):
            start_export(ctx, host, self.name, self._format, [], [], keys)
            return None
        query = [("host", host.key), ("action", self.name), ("format", self._format)]
        query += [("keys", host.record_key(r)) for r in records or []]
        ctx.redirect(ctx.url("_tw", "export") + "?" + urlencode(query))


# ---------------------------------------------------------------------- import
class ImportColumn:
    def __init__(self, name: str) -> None:
        self.name = name
        self._label: str | None = None
        self._required = False
        self._guesses: list[str] = []
        self._cast: Callable | None = None
        self._rules: list[Callable] = []
        self._relationship: tuple[str, str] | None = None
        self._example: Any = None
        self._numeric = False
        self._boolean = False

    @classmethod
    def make(cls, name: str) -> "ImportColumn":
        return cls(name)

    def label(self, label: str) -> "ImportColumn":
        self._label = label
        return self

    def required(self, condition: bool = True) -> "ImportColumn":
        self._required = condition
        return self

    def guess(self, headers: list[str]) -> "ImportColumn":
        """Other header names that should map to this column."""
        self._guesses = [h.lower() for h in headers]
        return self

    def cast_state_using(self, fn: Callable) -> "ImportColumn":
        self._cast = fn
        return self

    def rule(self, fn: Callable) -> "ImportColumn":
        """``fn(value)`` returns True/None when valid, or an error message."""
        self._rules.append(fn)
        return self

    def numeric(self) -> "ImportColumn":
        self._numeric = True
        return self

    def boolean(self) -> "ImportColumn":
        self._boolean = True
        return self

    def relationship(self, name: str, match_attribute: str) -> "ImportColumn":
        """Find the related record by ``match_attribute`` (e.g. category name)."""
        self._relationship = (name, match_attribute)
        return self

    def example(self, value: Any) -> "ImportColumn":
        self._example = value
        return self

    def get_label(self) -> str:
        return self._label or headline(self.name)

    def header_names(self) -> list[str]:
        return [self.name.lower(), self.get_label().lower(), *self._guesses]


class Importer:
    model: ClassVar[Any] = None
    columns: ClassVar[list[ImportColumn]] = []
    #: attribute used to find existing records to update (e.g. "email"); None = always create
    unique_by: ClassVar[str | None] = None

    @classmethod
    def get_columns(cls) -> list[ImportColumn]:
        return list(cls.columns)

    @classmethod
    def resolve_record(cls, ctx: "Context", data: dict[str, Any]) -> Any:
        from sqlalchemy import select

        if cls.unique_by and data.get(cls.unique_by) not in (None, ""):
            query = select(cls.model).where(getattr(cls.model, cls.unique_by) == data[cls.unique_by])
            # only the current tenant's records can be updated
            existing = ctx.db.scalars(ctx.panel.tenancy.scope(ctx, cls.model, query)).first()
            if existing is not None:
                return existing
        return cls.model()

    @classmethod
    def fill_record(cls, ctx: "Context", record: Any, data: dict[str, Any]) -> None:
        for key, value in data.items():
            setattr(record, key, value)

    @classmethod
    def before_save(cls, ctx: "Context", record: Any, data: dict[str, Any]) -> None:
        return None

    @classmethod
    def example_csv(cls) -> bytes:
        cols = cls.get_columns()
        return write_csv([[c.get_label() for c in cols], [c._example if c._example is not None else "" for c in cols]])


def read_table(content: bytes, filename: str) -> list[dict[str, str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        return [{headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)} for r in rows[1:]
                if any(v not in (None, "") for v in r)]
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in reader]


def run_import(ctx: "Context", importer: type[Importer], rows: list[dict[str, str]], tenancy_model: Any = None):
    """Import rows. Returns ``(created, updated, failures)``; failures are ``(row, message)``."""
    from sqlalchemy import func, select
    from sqlalchemy import inspect as sa_inspect

    created = updated = 0
    failures: list[tuple[dict, str]] = []
    columns = importer.get_columns()
    for row in rows:
        lookup = {k.lower(): v for k, v in row.items()}
        data: dict[str, Any] = {}
        errors = []
        for col in columns:
            raw = next((lookup[h] for h in col.header_names() if h in lookup), "")
            value: Any = raw.strip() if isinstance(raw, str) else raw
            if value in ("", None):
                if col._required:
                    errors.append(f"{col.get_label()} is required")
                continue
            try:
                if col._numeric:
                    value = Decimal(str(value).replace(",", ""))
                    value = int(value) if value == value.to_integral() else value
                if col._boolean:
                    value = str(value).lower() in ("1", "true", "yes", "y", "on")
                if col._relationship:
                    rel_name, attr = col._relationship
                    target = sa_inspect(importer.model).relationships[rel_name].mapper.class_
                    query = select(target).where(func.lower(getattr(target, attr)) == str(value).lower())
                    related = ctx.db.scalars(ctx.panel.tenancy.scope(ctx, target, query)).first()
                    if related is None:
                        raise ValueError(f"{col.get_label()} “{value}” was not found")
                    value = related
                if col._cast:
                    value = call(col._cast, state=value, row=row)
            except (InvalidOperation, ValueError) as exc:
                errors.append(str(exc) if str(exc) else f"{col.get_label()} is invalid")
                continue
            for rule in col._rules:
                result = call(rule, value=value, row=row)
                if result not in (True, None):
                    errors.append(result if isinstance(result, str) else f"{col.get_label()} is invalid")
            key = col._relationship[0] if col._relationship else col.name
            data[key] = value
        if errors:
            failures.append((row, "; ".join(errors)))
            continue
        try:
            with ctx.db.begin_nested():
                record = importer.resolve_record(ctx, data)
                is_new = sa_inspect(record).transient
                importer.fill_record(ctx, record, data)
                ctx.panel.tenancy.assign(ctx, importer.model, record)
                importer.before_save(ctx, record, data)
                ctx.db.add(record)
                ctx.db.flush()
            if is_new:
                created += 1
            else:
                updated += 1
        except Exception as exc:  # noqa: BLE001 - report the row, keep going
            failures.append((row, str(exc).splitlines()[0][:200]))
    ctx.db.commit()
    return created, updated, failures


class ImportAction(Action):
    def __init__(self, importer: type[Importer], name: str = "import") -> None:
        super().__init__(name)
        self.importer = importer
        self._label = "Import"
        self._icon = "upload"
        self._color = "gray"
        self._modal_heading = "Import records"
        self._modal_description = "Upload a CSV or Excel file. The first row must hold the column names."
        self._modal_submit_label = "Import"
        self._modal_icon = "upload"
        self._modal_width = "lg"
        self._form = True
        self._authorize = "create"
        self._style = "button"

    def build_form(self, ctx, host, record=None, records=None):  # type: ignore[override]
        from markupsafe import Markup

        from .forms import FileUpload, Placeholder
        from .forms.form import Form

        cols = ", ".join(c.get_label() + ("*" if c._required else "") for c in self.importer.get_columns())
        example = ctx.url("_tw", "import-example", host=host.key, action=self.name)
        link = Markup('<a href="{}" class="text-sm font-medium text-primary-600 hover:text-primary-500" '
                      'hx-boost="false" download>{}</a>').format(example, __("Download example CSV"))
        return Form().columns(1).schema([
            FileUpload("file").label("File").accepted_file_types([".csv", ".xlsx", "text/csv"])
            .directory("imports").max_size(10240).required(),
            Placeholder("columns").label("Expected columns").content(cols),
            Placeholder("example").hidden_label().content(link),
        ])

    def run(self, ctx, host, record=None, records=None, data=None, form=None):  # type: ignore[override]
        from .notifications import Notification

        storage = ctx.panel.storage
        path = (data or {}).get("file")
        content = storage.path(path).read_bytes()
        rows = read_table(content, path)
        created, updated, failures = run_import(ctx, self.importer, rows)
        storage.delete(path)
        note = Notification("Import finished").body(
            __(":created created, :updated updated, :failed failed.", created=created, updated=updated,
               failed=len(failures)) if failures
            else __(":created created, :updated updated.", created=created, updated=updated))
        if failures:
            headers = list(rows[0].keys()) if rows else []
            report = write_csv([headers + ["Error"]] + [[r.get(h, "") for h in headers] + [msg] for r, msg in failures])
            stored = storage.save(io.BytesIO(report), "failed-rows.csv", "imports")
            note.warning().persistent().action("Download failed rows", storage.url(stored))
        else:
            note.success()
        note.send(ctx)
